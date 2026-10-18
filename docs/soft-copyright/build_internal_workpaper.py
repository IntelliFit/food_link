from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "智健食探软著申报事实核对与人工撰写工作底稿_V3.0.2.docx"

NAVY = "17324D"
BLUE = "2E74B5"
TEAL = "00856A"
INK = "1F2937"
MUTED = "5B6573"
LIGHT = "F2F4F7"
PALE_BLUE = "E8F0F7"
PALE_TEAL = "E8F6F2"
PALE_RED = "FDECEC"
RED = "9B1C1C"
WHITE = "FFFFFF"
BORDER = "CBD5E1"


def set_run_font(run, name="Microsoft YaHei", size=None, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[min(idx, len(widths_dxa) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])


def add_text(doc, text, *, bold=False, italic=False, color=INK, size=10.5, align=None,
             before=0, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullet(doc, text, level=0, bold_prefix=None):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=10.5, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_callout(doc, title, body, fill=PALE_BLUE, title_color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.5, color=title_color, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.12
    r2 = p2.add_run(body)
    set_run_font(r2, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_kv_table(doc, rows, widths=(2700, 6660)):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0]
    hdr.cells[0].text = "项目"
    hdr.cells[1].text = "核验结果 / 待办"
    set_repeat_table_header(hdr)
    for cell in hdr.cells:
        set_cell_shading(cell, NAVY)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                set_run_font(run, size=9.5, color=WHITE, bold=True)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        set_cell_shading(cells[0], LIGHT)
        for run in cells[0].paragraphs[0].runs:
            set_run_font(run, size=9.5, color=NAVY, bold=True)
        for run in cells[1].paragraphs[0].runs:
            set_run_font(run, size=9.5, color=INK)
        for p in cells[0].paragraphs + cells[1].paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
    set_table_geometry(table, list(widths))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=8.7, color=WHITE, bold=True)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            if len(table.rows) % 2 == 0:
                set_cell_shading(cells[i], "FAFBFC")
            for run in cells[i].paragraphs[0].runs:
                set_run_font(run, size=8.5, color=INK)
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
            cells[i].paragraphs[0].paragraph_format.line_spacing = 1.0
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_source(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    r1 = p.add_run(label + "：")
    set_run_font(r1, size=8.5, color=MUTED, bold=True)
    r2 = p.add_run(url)
    set_run_font(r2, size=8.5, color=BLUE)
    return p


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.25)
    section.bottom_margin = Cm(2.15)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, NAVY, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("智健食探软著申报工作底稿")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    r2 = p.add_run("    内部资料 · AI 辅助整理 · 不得直接提交")
    set_run_font(r2, size=8.5, color=RED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = fp.add_run("内部事实核对与人工撰写工作底稿  |  第 ")
    set_run_font(r3, size=8, color=MUTED)
    add_page_field(fp)
    r4 = fp.add_run(" 页")
    set_run_font(r4, size=8, color=MUTED)


def build():
    doc = Document()
    setup_styles(doc)

    # First-page masthead
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("SOFTWARE COPYRIGHT WORKPAPER")
    set_run_font(r, size=9, color=TEAL, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    r = title.add_run("智健食探软著申报事实核对\n与人工撰写工作底稿")
    set_run_font(r, size=23, color=NAVY, bold=True)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(14)
    r = sub.add_run("拟申报软件：智健食探智能营养健康管理软件 V3.0.2")
    set_run_font(r, size=12, color=MUTED, bold=True)

    add_kv_table(doc, [
        ("底稿用途", "供真实研发人员、权利人和经办人核验事实、确定申报边界并组织人工撰写；不是登记申请材料。"),
        ("项目基线", "food_link 仓库，dev 分支；根 package.json 版本 3.0.2；核对日期 2026-08-25。"),
        ("当前状态", "工作区存在大量未提交改动，尚未形成可直接申报的固定版本。必须先冻结 commit/tag。"),
        ("待确认主体", "著作权人、开发方式、权利取得方式、完成日期、首次发表日期、经办人及身份证明均待权利人核验。"),
    ])

    add_callout(
        doc,
        "关键合规结论（请先读）",
        "2026-03-15 启用的新版申请表要求经办人实名声明未使用 AI 开发编写代码、撰写文档或生成登记申请材料，并加入失信责任。本文由 AI 辅助整理，不能复制、改名或直接上传。当前项目也存在 AI 编码代理参与开发的工作记录；在获得中国版权保护中心对“AI 辅助、人主导复核/修改”的明确口径前，不建议为当前 V3.0.2 签署该声明。",
        fill=PALE_RED,
        title_color=RED,
    )
    add_text(doc, "编制日期：2026 年 8 月 25 日", color=MUTED, size=9, after=0)

    doc.add_page_break()

    add_heading(doc, "1. 结论与建议", 1)
    add_callout(
        doc,
        "建议结论",
        "先做“申报资格与版本冻结”，再做材料。若经办人不能据实签署新版 AI 声明，就不要递交；不得通过人工抄写、代办或删除 AI 痕迹来规避。可以先向中国版权保护中心咨询并保留工单/通话记录，或选择一个能够真实满足声明的历史版本。",
        fill=PALE_TEAL,
        title_color=TEAL,
    )
    for item in [
        "软件边界建议：以用户端微信小程序 + 专属 Go 后端的智能营养健康管理能力为一个软件整体；排除 Admin 管理后台、E2E 测试、数据爬取、压测和部署脚本。",
        "名称建议：智健食探智能营养健康管理软件；简称“智健食探”；版本号建议与申报 commit 对应并写为 V3.0.2。最终名称必须与申请表、说明书、截图、源代码页眉一致。",
        "当前不具备直接申报条件：dev 工作区有未提交改动，申请人及权属信息缺失，且 AI 使用声明存在实质风险。",
        "本底稿可用于内部尽调、研发访谈、法务核验和向登记机构咨询；不生成可直接签署的 500–1300 字主要功能正文，也不生成源代码/说明书鉴别材料。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2. 2026 年最新规则核对", 1)
    add_heading(doc, "2.1 现行法定基础", 2)
    add_text(doc, "《计算机软件著作权登记办法》目前仍是软件登记的法定基础。与本项目直接相关的要求如下：")
    rules = [
        ("申请主体", "申请人应当是软件著作权人，或通过继承、受让、承受取得权利的自然人、法人或其他组织。"),
        ("软件条件", "软件应为独立开发，或经原著作权人许可、在功能或性能方面有重要改进。"),
        ("三类材料", "申请表、软件鉴别材料、相关证明文件。"),
        ("程序鉴别材料", "源程序前、后各连续 30 页；不足 60 页提交全部；除特定情况外每页不少于 50 行。"),
        ("文档鉴别材料", "任一种文档前、后各连续 30 页；不足 60 页提交全部；除特定情况外每页不少于 30 行。"),
        ("证明文件", "身份证明；存在合同、任务书、原软件许可、继承或受让情形时提交相应权属证明。"),
        ("形式", "统一表格、中文填写；外文证明附中文译本；文件使用 A4 纸。"),
        ("审查与补正", "法定审查期限为受理日起 60 日；收到补正通知后 30 日内补正，逾期视为撤回。"),
    ]
    add_matrix(doc, ["事项", "法定要求摘要"], rules, [2400, 6960])

    add_heading(doc, "2.2 2026-03-15 新版申请表变化", 2)
    changes = [
        ("主要功能字数", "由旧表 200 字以内调整为 500–1300 字；应具体说明研发背景、架构、功能、技术实现和应用场景。"),
        ("诚信制度", "建立软件版权登记申请诚信制度，引入失信惩戒机制。"),
        ("AI 声明", "签章页增加“未使用 AI 开发编写代码、撰写文档或生成登记申请材料”的声明。"),
        ("经办人实名", "经办人签名并填写身份证号码；公开转发材料显示承诺内容需本人手写，正式操作以系统生成的新表为准。"),
        ("新表生效", "优化后的申请表自 2026-03-15 发布之日起施行；不要保存或套用旧版离线表。"),
    ]
    add_matrix(doc, ["变化", "对本项目的影响"], changes, [2200, 7160])
    add_callout(
        doc,
        "证据强度说明",
        "AI 声明和 500–1300 字要求来自中国版权保护中心 2026-03-15 通知及新版申请表的公开转发/新闻核验；国家版权局网站公开的《登记办法》仍提供基础法定要求。提交当天必须再次登录登记系统核对最新表单字段与提示。",
        fill=PALE_BLUE,
    )

    add_heading(doc, "2.3 实际办理口径", 2)
    for item in [
        "通过中国版权保护中心著作权登记系统注册、实名认证、在线填报，并使用系统当次生成的申请表。不要从代理网站下载旧表。",
        "申请人可自行办理或委托代理，但权利人和经办人仍需对真实性负责；代理不能替代研发人员确认功能、代码和权属。",
        "普通登记的实际处理时长、是否收费、证书形态及上传文件格式，以提交当日系统页面为准；法定规则中的 60 日从受理日起计算。",
        "保留一整套提交副本和版本 hash，便于补正时保持一致。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3. 申报资格与 AI 使用风险", 1)
    add_heading(doc, "3.1 当前风险判断", 2)
    risk_rows = [
        ("登记材料由 AI 生成", "已发生", "本工作底稿即为 AI 辅助产物，不能直接转为申请表正文或说明书。", "红"),
        ("项目代码存在 AI 辅助开发", "高概率/需审计", "项目由编码代理长期参与，需由负责人逐 commit 审计并向登记机构咨询。", "红"),
        ("申报版本未冻结", "已发生", "dev 工作区存在大量未提交修改，截图、源码和版本无法形成单一证据链。", "红"),
        ("著作权人未确认", "未完成", "package.json author 为空；多人提交不等于权属已明确。需核对劳动/合作/委托协议。", "黄"),
        ("第三方/开源成分", "存在", "Taro、React、Gin、GORM 等依赖不得作为自研代码主张；应保留许可证清单。", "黄"),
        ("软件边界过宽", "可控", "若把微信端、App、Admin、测试、爬虫一起申报，名称和鉴别材料容易不一致。", "黄"),
    ]
    add_matrix(doc, ["风险", "状态", "判断", "等级"], risk_rows, [1900, 1350, 4910, 1200])

    add_heading(doc, "3.2 必须由权利人作出的决定", 2)
    decisions = [
        "确认申报版本是否完全没有使用 AI 编写代码、文档或申请材料；不要把“人工看过”自动等同于满足声明。",
        "以事实向中国版权保护中心咨询：AI 只做补全/建议、人工逐行修改是否仍属于“使用 AI 开发编写代码”；AI 仅参与内部底稿但最终由研发人员独立重写是否可以签署。",
        "若当前版本不满足声明，选择可真实满足的历史 commit，或暂停申报，等待明确政策口径；不得倒签、隐藏或伪造开发记录。",
        "确定著作权人：个人、公司、共同著作权人或受让主体；核对开发者与权利人的劳动、委托、合作、转让关系。",
    ]
    for item in decisions:
        add_number(doc, item)

    add_heading(doc, "3.3 咨询登记机构时建议逐字说明的事实", 2)
    add_callout(
        doc,
        "咨询问题模板（事实问题，不是申报文本）",
        "“我们有一款已持续开发的软件，部分代码提交和内部技术文档曾使用生成式 AI 辅助，但均由研发人员审查、修改并承担责任。2026 新版申请表要求声明未使用 AI 开发编写代码、撰写文档或生成登记申请材料。请问：这种人主导、AI 辅助的情况能否申请？若不能，能否申报未使用 AI 的历史版本？需要保留或提交哪些证明？”",
        fill=PALE_TEAL,
        title_color=TEAL,
    )

    add_heading(doc, "4. 推荐的软件申报边界", 1)
    add_heading(doc, "4.1 名称和范围", 2)
    add_kv_table(doc, [
        ("建议全称", "智健食探智能营养健康管理软件"),
        ("建议简称", "智健食探"),
        ("建议版本", "V3.0.2（须与冻结 commit 对应；若改用历史版本，版本号同步调整）"),
        ("软件类型", "应用软件 / 健康管理工具（最终按系统分类项选择）"),
        ("纳入范围", "微信小程序客户端、专属业务 API、分析任务与营养健康数据处理逻辑。"),
        ("排除范围", "Admin 后台、测试后台、E2E/压测、校园数据爬取、部署脚本、第三方依赖、构建产物。"),
        ("不主张内容", "第三方模型能力、开源框架本身、第三方数据源、微信平台能力及云服务本身。"),
    ])

    add_heading(doc, "4.2 技术架构事实", 2)
    add_text(doc, "用户端以 Taro 4.1.10、React 18、TypeScript 5.4 和 Sass 构建微信小程序；服务端为 Go 1.26 + Gin/GORM，使用 PostgreSQL 持久化、Redis/Kafka 支持缓存与异步处理、腾讯云 COS/CDN 管理用户图片，并通过受控外部模型接口完成图像/文本分析。")
    add_text(doc, "逻辑链路：微信小程序 → HTTPS API → 业务服务（分析、记录、健康、校园、社区等）→ PostgreSQL/Redis/Kafka → COS/CDN 与外部 AI 服务。", bold=True, color=NAVY)
    add_callout(
        doc,
        "医疗边界",
        "软件定位是饮食记录、营养估算和健康管理辅助，不应在申请材料中描述为疾病诊断、治疗或替代医生。涉及健康档案和图片时，应如实说明授权、访问控制和隐私保护。",
        fill=PALE_BLUE,
    )

    add_heading(doc, "5. 申请表事实采集单", 1)
    add_text(doc, "下表仅列出事实项。方括号内容必须由申请人核验，不得推测填写。")
    add_kv_table(doc, [
        ("软件全称", "建议：智健食探智能营养健康管理软件"),
        ("简称", "建议：智健食探"),
        ("版本号", "建议：V3.0.2；以冻结 commit 的 package.json 为准"),
        ("著作权人", "[待确认：个人/公司全称，必须与证件一致]"),
        ("统一社会信用代码/身份证", "[待确认，不在内部文档传播完整号码]"),
        ("开发方式", "[待确认：独立开发/合作开发/委托开发/下达任务开发]"),
        ("权利取得方式", "[待确认：原始取得/继受取得；如继受需证明]"),
        ("权利范围", "[待确认：全部权利/部分权利]"),
        ("开发完成日期", "[待确认：对应申报版本实际完成日期，不得用底稿日期代替]"),
        ("首次发表状态", "[待确认：未发表/已发表；如已发表，填写真实日期与地点]"),
        ("开发硬件环境", "常规 x86_64/AMD64 开发计算机与云服务器；具体 CPU、内存按实际填写"),
        ("运行硬件环境", "支持微信的移动终端 + Linux AMD64 服务端；云资源规格按冻结版本实际环境填写"),
        ("开发软件环境", "Windows/macOS/Linux 开发环境；Node.js、Taro 4.1.10、TypeScript 5.4、Go 1.26 等，按实际开发机补全"),
        ("运行软件环境", "微信客户端/微信小程序运行环境；Linux、PostgreSQL、Redis/Kafka 等服务，按申报边界核实"),
        ("编程语言", "TypeScript/TSX、Go、SCSS；是否列 SQL 等由冻结源码范围决定"),
        ("源程序量", "当前粗略核对：微信端 TS/TSX 约 79,799 行、SCSS 约 60,695 行；后端非测试 Go 约 150,580 行。正式值须在冻结 commit 重新统计。"),
        ("经办人", "[待确认：申请人内部真实经办人；需能据实签署新版承诺]"),
    ])

    add_heading(doc, "6. 项目功能事实库（供研发人员独立撰写）", 1)
    add_callout(
        doc,
        "使用方式",
        "研发人员应只把本节当作核对清单，回到冻结版本亲自操作、截图和描述。不得把本节改写、拼接或复制到申请表的“主要功能”栏。",
        fill=PALE_RED,
        title_color=RED,
    )
    modules = [
        ("登录与账户", "微信登录、手机号快捷登录、用户资料、隐私与账号安全", "app.config；login/profile 页面；user API"),
        ("AI 食物分析", "拍照、文字、批量/多图、精准多角度、视频关键帧等识别入口；异步任务与结果回看", "analyze、analyze-loading、result 页面；analyze service"),
        ("饮食记录", "分析结果保存、手动记录、文本记录、餐次管理、历史记录与详情", "record-* 页面；foodrecord service"),
        ("营养与趋势", "每日摄入总览、宏量/微量营养、体重、饮水、运动、身体指标趋势", "首页、stats、body-* 页面；stats service"),
        ("健康档案", "个人健康资料、目标管理、健康报告图片上传与结构化提取", "health-profile 页面；health API"),
        ("校园食堂", "学校/校区/食堂/楼层目录、校园菜品检索与分享、校园餐推荐", "campus-* 页面；publicfood/health service"),
        ("公共食物与社区", "食物库、收藏、反馈、动态、评论、点赞、互动通知与好友", "food-library/community 页面；publicfood/community API"),
        ("菜谱", "个人菜谱创建、编辑、查看和按菜谱生成记录", "recipes 页面与 API"),
        ("保质期管理", "食品录入、到期看板、状态更新与订阅提醒", "expiry 页面；expiry service"),
        ("补剂管理", "公共补剂库、标签图片识别、补剂记录与每日计划", "supplements 页面；supplement service"),
        ("宠物/对话助手", "宠物形象、饮食对话、基于真实用户目标与校园食物数据的工具调用推荐", "pet-home/pet-chat 页面；pet/health agent service"),
        ("会员与激励", "积分、签到、奖励中心及会员权益控制；支付属于平台合规独立事项", "pro-membership/reward 页面；membership API"),
    ]
    add_matrix(doc, ["模块", "可核验功能事实", "代码/页面证据"], modules, [1600, 4800, 2960])

    add_heading(doc, "7. 人工撰写“主要功能”访谈提纲", 1)
    add_text(doc, "新版主要功能栏要求 500–1300 字。请由了解产品和代码的真实研发人员，在不查看任何 AI 成稿的情况下独立回答以下问题，再由法务核对事实和一致性。")
    prompts = [
        "研发背景：用户在饮食记录、营养理解和持续体重管理上遇到的具体问题是什么？为什么要做微信小程序？",
        "目标用户与场景：面向哪些用户？在家庭餐、外卖、校园食堂、包装食品、补记录等场景如何使用？",
        "软件边界：本次申请包含哪些客户端和服务端能力？明确排除哪些后台、工具和第三方能力？",
        "核心架构：客户端、API、异步任务、数据库、对象存储、缓存/消息和外部模型如何协作？",
        "核心功能：选 5–8 个已在冻结版本真实可用、能由截图和代码同时证明的模块逐项说明。",
        "关键流程：从用户选择图片/输入文字，到上传、提交任务、轮询/提醒、展示结果、保存记录的完整流程是什么？",
        "技术特点：哪些数据校验、任务幂等、缓存策略、营养匹配、权限控制或隐私保护属于自研实现？",
        "原创性边界：自研代码与 Taro、React、Gin、GORM、微信接口、外部模型和第三方数据如何区分？",
        "应用价值：软件怎样降低记录门槛、形成摄入—趋势—行动闭环？不要使用无法证明的疗效或准确率。",
        "版本事实：上述每项是否都能在同一冻结 commit 中运行、截图并找到代码？",
    ]
    for item in prompts:
        add_number(doc, item)

    add_heading(doc, "8. 人工编制说明书的结构与截图清单", 1)
    add_heading(doc, "8.1 推荐文档类型", 2)
    add_text(doc, "建议提交《智健食探智能营养健康管理软件 V3.0.2 用户操作说明书》。用户手册比纯设计说明书更容易用真实截图证明功能，但必须由研发/产品人员基于冻结版本人工编制。")
    add_heading(doc, "8.2 章节结构", 2)
    chapters = [
        "封面：全称、版本号、著作权人、完成日期（均与申请表一致）。",
        "软件概述：研发目的、适用对象、功能边界、免责声明。",
        "运行环境：微信版本/基础库、网络要求、服务端环境（只写实际）。",
        "登录与首次使用：授权、手机号、协议、健康资料和目标设置。",
        "食物分析：拍照/选图、文字、精准模式、任务等待、结果确认与保存。",
        "饮食记录与首页：日总览、餐次、记录详情、历史记录。",
        "营养统计与身体数据：趋势、体重、饮水、运动和身体指标。",
        "健康档案：资料查看/修改、报告上传与隐私说明。",
        "校园、食物库与社区：按申报版本真实页面择要说明。",
        "保质期、补剂、菜谱与对话助手：只纳入冻结版本可稳定复现的功能。",
        "数据与安全：账号权限、上传图片、缓存清理、注销与隐私入口。",
        "常见问题：网络失败、识别失败、任务等待、缓存、权限与错误提示。",
    ]
    for item in chapters:
        add_number(doc, item)

    add_heading(doc, "8.3 最小截图证据集", 2)
    shots = [
        ("登录/首页", "登录页、首页日总览、中央识别入口", "显示软件品牌和真实数据状态；隐去手机号/头像等个人信息"),
        ("分析流程", "选择图片、分析方式、等待页、结果页、保存成功", "同一测试账号和同一版本；不要拼接不同版本 UI"),
        ("记录与统计", "记录列表/详情、营养趋势、体重/饮水/运动", "示例数据应合法、脱敏、可复现"),
        ("健康档案", "资料页、目标设置、报告入口", "不得展示真实健康报告或身份证明"),
        ("校园/食物库", "食堂筛选、菜品详情、收藏/反馈", "仅使用当前版本真实上线功能"),
        ("扩展模块", "保质期、补剂、菜谱、宠物对话", "根据申报范围择要，不求功能数量"),
        ("隐私与安全", "隐私政策、账号安全、清除缓存、注销入口", "体现数据处理边界，不声称不实际具备的加密措施"),
    ]
    add_matrix(doc, ["场景", "建议截图", "注意事项"], shots, [1700, 3600, 4060])

    add_heading(doc, "9. 源程序鉴别材料的人工作业方案", 1)
    add_callout(
        doc,
        "本底稿不生成源码交存件",
        "新版申请表包含未使用 AI 生成登记申请材料的声明；自动拼接、排版或补齐源代码页可能成为申报材料生成行为。以下只提供人工操作的控制点。",
        fill=PALE_RED,
        title_color=RED,
    )
    steps = [
        "冻结版本：清理无关改动后，由负责人确定 commit hash，并创建只读归档；记录 package.json 版本和构建时间。",
        "确定自研边界：只选申请人拥有著作权的 TypeScript/TSX 与 Go 业务代码；排除 node_modules、vendor、dist、测试、生成代码、迁移数据、密钥和第三方源码。",
        "建立人工文件顺序：按客户端入口→核心业务→服务端入口→核心服务的稳定顺序登记文件清单；不得为凑页数重复代码。",
        "核对连续性：按该顺序形成完整源程序，取前 30 页和后 30 页；不足 60 页则全部提交。每页除最后一页外保持不少于 50 行。",
        "统一页眉页码：每页写同一软件全称和版本号，右上角连续页码；名称不得与申请表或说明书不同。",
        "保密检查：不提交 token、密码、OAuth code、数据库连接串、真实手机号、健康数据、用户图片、模型密钥或商业敏感配置。必要时按法定例外交存方式办理，不要随意涂黑。",
        "权属检查：排除开源框架、复制片段、外包方未转让代码和权属不清的 AI 生成片段；保留许可证及开发记录。",
        "双人复核：研发负责人核对代码真实性和可编译性，法务/经办人核对名称、版本、页数、行数、权属和敏感信息。",
    ]
    for item in steps:
        add_number(doc, item)

    add_heading(doc, "9.1 当前可核验的源码规模", 2)
    add_kv_table(doc, [
        ("微信端 TypeScript/TSX", "约 289 个文件、79,799 行（当前工作区粗略统计）"),
        ("微信端 SCSS", "约 107 个文件、60,695 行（当前工作区粗略统计；是否作为程序鉴别材料需人工决定）"),
        ("Go 后端非测试代码", "约 462 个文件、150,580 行（排除 *_test.go，当前工作区粗略统计）"),
        ("统计限制", "工作区有未提交改动，且统计未做生成代码/第三方片段逐文件权属过滤，不能直接填写申请表。"),
    ])

    add_heading(doc, "9.2 候选业务文件（仅供人工定位）", 2)
    candidate_rows = [
        ("首页与日总览", "apps/wechat/src/pages/index/index.tsx", "约 4,025 行"),
        ("图片分析入口", "apps/wechat/src/packageExtra/pages/analyze/index.tsx", "约 1,928 行"),
        ("分析结果", "apps/wechat/src/packageExtra/pages/result/index.tsx", "约 3,582 行"),
        ("统计页", "apps/wechat/src/pages/stats/index.tsx", "约 1,952 行"),
        ("健康档案", "apps/wechat/src/packageExtra/pages/health-profile/index.tsx", "约 1,337 行"),
        ("校园食堂", "apps/wechat/src/packageExtra/pages/campus-canteen/index.tsx", "约 1,140 行"),
        ("分析服务", "backend/internal/analyze/service/analyze_service.go", "约 8,185 行"),
        ("统计服务", "backend/internal/health/service/stats_service.go", "约 2,801 行"),
        ("校园餐 Agent", "backend/internal/health/service/campus_diet_agent_service.go", "约 1,513 行"),
        ("宠物服务", "backend/internal/pet/service/pet_service.go", "约 1,251 行"),
    ]
    add_matrix(doc, ["模块", "文件", "当前行数"], candidate_rows, [1700, 5860, 1800])

    add_heading(doc, "10. 身份与权属证明清单", 1)
    proofs = [
        ("企业申请", "营业执照/主体证明；经办人信息与授权（按系统要求）；公章签章页", "单位全称、证件号、地址一致"),
        ("个人申请", "有效身份证明；本人签名及系统要求的信息", "姓名和号码与实名认证一致"),
        ("职务开发", "劳动合同、岗位/任务书、职务成果制度或权属条款", "确认软件著作权归属用人单位的事实基础"),
        ("合作开发", "合作开发协议、贡献和权利分配、代表申请授权", "共同著作权人及署名一致"),
        ("委托开发", "委托合同及明确的软件著作权归属条款", "没有明确约定时风险较高"),
        ("受让/继承", "转让、继承或承受证明", "权利链连续"),
        ("开源组件", "依赖清单、许可证文本、自研与第三方边界说明", "不把第三方代码作为自研交存"),
        ("开发证据", "Git 记录、需求/设计记录、测试报告、发布记录、构建 hash", "与冻结版本一致并妥善留存"),
        ("AI 使用审计", "AI 工具使用范围、相关 commit、人工修改记录、官方咨询结论", "决定是否能够据实签署新版声明"),
    ]
    add_matrix(doc, ["情形", "应准备文件", "核验点"], proofs, [1700, 4300, 3360])

    add_heading(doc, "11. 在线办理流程与内部门禁", 1)
    process = [
        "资格门禁：确认能够据实签署最新版申请表全部声明，尤其是 AI 使用声明；否则停止。",
        "权属门禁：确定著作权人、开发方式、取得方式和权利范围，补齐合同/任务书/转让证明。",
        "版本门禁：冻结 commit/tag，确认名称、V3.0.2、完成日期和发表状态。",
        "人工材料：研发人员独立撰写主要功能和用户手册，人工整理自研源代码鉴别材料。",
        "一致性复核：申请表、说明书、截图、源码页眉、版本和权利人逐项比对。",
        "敏感信息复核：清理密钥、用户数据、健康资料、内部域名和不必要的商业机密。",
        "登录中国版权保护中心系统：实名认证、按当次表单在线填报、上传系统要求文件。",
        "签章：下载/打印系统生成的新表，由真实经办人按页面要求本人手写、签名、填写身份信息；企业按要求盖章。",
        "提交与留档：保存完整提交副本、文件 hash、提交时间和受理号。",
        "补正：收到通知后在 30 日内由同一负责人核对并补正，保持版本和内容一致。",
    ]
    for item in process:
        add_number(doc, item)

    add_heading(doc, "12. 一致性复核矩阵", 1)
    checks = [
        ("软件全称", "申请表、说明书封面/页眉、源码页眉、截图标题", "完全一致"),
        ("版本号", "package.json、冻结 tag、申请表、说明书、源码页眉", "完全一致；V 大小写统一"),
        ("权利人", "申请表、身份证明、合同/任务书、说明书署名", "法定名称一致"),
        ("完成/发表日期", "Git/发布记录、申请表、说明书", "真实且逻辑顺序正确"),
        ("主要功能", "申请表、说明书章节、实际 UI、源码", "每项有同版本证据"),
        ("技术环境", "申请表、说明书、依赖版本、部署事实", "不写过期 Python 后端或未纳入范围的 App"),
        ("源程序量", "冻结版本统计报告、申请表", "统计口径可复算"),
        ("AI 使用", "新版声明、开发记录、咨询结论", "只能如实，不得形式规避"),
        ("隐私与健康", "说明书、隐私政策、实际权限", "不披露真实数据，不夸大医疗能力"),
    ]
    add_matrix(doc, ["项目", "对照材料", "通过标准"], checks, [1800, 4700, 2860])

    add_heading(doc, "13. 常见补正/不予登记风险", 1)
    for item in [
        "软件名称、版本号或权利人署名在不同材料中不一致。",
        "主要功能仍是模板化套话，与截图和代码无法对应，或把未上线规划写成现有功能。",
        "程序/文档页数、连续性或每页行数不符合要求；用截图、配置、空行或重复代码凑页。",
        "提交的不是软件程序和文档，或包含大量第三方、开源、生成代码而无法说明权属。",
        "合作/委托/职务开发未提供足够权属文件，多位开发者与单一申请人之间缺少权利链。",
        "申请表使用旧版，或经办人未按新表要求本人完成承诺、签名和身份信息。",
        "对 AI 使用情况作出与开发记录不符的声明。",
        "补正超过 30 日，或补正时又换了版本、截图和源代码。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "14. 建议的下一步", 1)
    next_steps = [
        "由项目负责人和法务确定著作权申请主体，并审查劳动/合作/委托/转让文件。",
        "对全仓 Git 历史和开发流程做 AI 使用审计，明确哪些版本、哪些文件使用过 AI。",
        "携带真实事实联系中国版权保护中心，确认新版 AI 声明对人机协作代码和人工重写材料的适用口径，并留存答复。",
        "只有在可以据实签署时，冻结一个 commit/tag；当前有未提交改动，不能直接使用工作目录。",
        "由真实研发人员在冻结版本上独立撰写主要功能和用户手册、人工整理代码交存件；法务只做事实与一致性核验。",
        "提交当天重新核对登记系统的新表字段、签章页和上传要求。",
    ]
    for item in next_steps:
        add_number(doc, item)

    doc.add_page_break()
    add_heading(doc, "附录 A：事实来源与法规链接", 1)
    add_source(doc, "国家版权局《计算机软件著作权登记办法》", "https://www.ncac.gov.cn/xxfb/flfg/bmgz/202410/t20241015_869486.html")
    add_source(doc, "深圳市软件行业协会转发中国版权保护中心 2026-03-15 通知及新旧表对比", "https://www.ssia.org.cn/page63?article_id=2588")
    add_source(doc, "澎湃新闻对 2026 新版申请表与经办人实名要求的核验报道", "https://m.thepaper.cn/newsDetail_forward_32781466")
    add_source(doc, "中国版权保护中心著作权登记系统入口", "https://register.ccopyright.com.cn/")
    add_text(doc, "项目内部事实来源", bold=True, color=NAVY, size=10.5, before=8)
    for item in [
        "package.json（版本 3.0.2、Taro/React 等依赖）",
        "apps/wechat/src/app.config.ts（页面、分包和功能入口）",
        "docs/食探-商业计划书-精简版.md（产品定位）",
        "backend/go.mod（Go 版本和服务端依赖）",
        "前后端业务源码与 2026-08-25 本地行数统计",
        "项目持久化状态文件（品牌名、当前任务和技术决策）",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "附录 B：申报前签字检查页", 1)
    add_text(doc, "以下项目均由申请人内部真实负责人填写。本页是内部控制记录，不提交登记机构。")
    sign_rows = [
        ("申报软件全称/版本", "____________________________________________"),
        ("冻结 commit/tag", "____________________________________________"),
        ("著作权人", "____________________________________________"),
        ("开发方式/取得方式", "____________________________________________"),
        ("AI 使用审计结论", "____________________________________________"),
        ("中国版权保护中心咨询记录编号/日期", "____________________________________________"),
        ("研发负责人复核", "签名：________________  日期：____________"),
        ("法务/权属负责人复核", "签名：________________  日期：____________"),
        ("经办人确认可据实签署最新版申请表", "签名：________________  日期：____________"),
    ]
    add_kv_table(doc, sign_rows)
    add_callout(
        doc,
        "最终提醒",
        "只有当申请人能够对最新版申请表中的全部声明承担真实、完整的法律责任时，才进入提交。任何人不得把本 AI 辅助底稿当作“未使用 AI 生成登记申请材料”的证明。",
        fill=PALE_RED,
        title_color=RED,
    )

    props = doc.core_properties
    props.title = "智健食探软著申报事实核对与人工撰写工作底稿 V3.0.2"
    props.subject = "软件著作权申报内部合规工作底稿"
    props.author = "food_link 项目组（AI 辅助整理，须人工核验）"
    props.keywords = "智健食探, 软件著作权, 内部底稿, AI合规"
    props.comments = "本文件不得直接作为软件著作权登记申请材料提交。"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
