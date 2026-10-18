from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


path = Path(__file__).resolve().parent / "智健食探软著申报事实核对与人工撰写工作底稿_V3.0.2.docx"
doc = Document(path)
errors = []

section = doc.sections[0]
if not (11800 <= section.page_width.twips <= 12020 and 16750 <= section.page_height.twips <= 16920):
    errors.append(f"page_geometry={section.page_width.twips}x{section.page_height.twips}")

headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading ")]
if len(headings) < 20:
    errors.append(f"too_few_headings={len(headings)}")

text = "\n".join(p.text for p in doc.paragraphs)
for table in doc.tables:
    for row in table.rows:
        text += "\n" + "\t".join(cell.text for cell in row.cells)

required = [
    "内部资料 · AI 辅助整理 · 不得直接提交",
    "未使用 AI 开发编写代码、撰写文档或生成登记申请材料",
    "500–1300 字",
    "智健食探智能营养健康管理软件",
    "中国版权保护中心",
]
for item in required:
    if item not in text and item not in section.header.paragraphs[0].text:
        errors.append(f"missing_required_text={item}")

for token in ("turn0search", "turn1search", "cite", ":codex-file-citation"):
    if token in text:
        errors.append(f"internal_citation_token={token}")

for idx, table in enumerate(doc.tables, 1):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    grid_cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
    if tbl_w is None or tbl_w.get(qn("w:type")) != "dxa":
        errors.append(f"table_{idx}_missing_dxa_width")
    if tbl_ind is None or tbl_ind.get(qn("w:type")) != "dxa":
        errors.append(f"table_{idx}_missing_dxa_indent")
    if not grid_cols:
        errors.append(f"table_{idx}_missing_grid")
    for r_idx, row in enumerate(table.rows, 1):
        for c_idx, cell in enumerate(row.cells, 1):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None or tc_w.get(qn("w:type")) != "dxa":
                errors.append(f"table_{idx}_r{r_idx}_c{c_idx}_missing_dxa_width")

print(f"file={path}")
print(f"paragraphs={len(doc.paragraphs)} headings={len(headings)} tables={len(doc.tables)}")
print(f"page_twips={section.page_width.twips}x{section.page_height.twips}")
print(f"characters={len(text)}")
print(f"errors={len(errors)}")
for error in errors:
    print(error)

raise SystemExit(1 if errors else 0)
